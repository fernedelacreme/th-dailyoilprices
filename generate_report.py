import calendar
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np


def build_monthly_report(raw_df: pd.DataFrame, target_year: int, target_month: int, output_path: str):
    """
    ฟังก์ชันแปลงข้อมูลราคาน้ำมันดิบให้ออกมาเป็นรายงาน Excel 3 ชีท
    - 7 หมวดหลักแสดงสถานีบริการตามลิสต์เป๊ะๆ พร้อมกลไก Proxy/Fallback แหล่งราคา
    - หมวด 2: ดีเซลหมุนเร็ว (HSD, 0.005%S) ดึงจาก 'ดีเซลหมุนเร็ว HSD B7'
    - หมวดอื่นๆ แสดงต่อท้ายด้านล่าง และใส่ครบทุกสถานีบริการที่มีในไฟล์ raw
    - ช่องที่ไม่มีราคาให้ใส่ 0.00
    - ชีท 3 ไฮไลต์ส่วนต่างรายวัน แดง = เพิ่มขึ้น, เขียว = ลดลง
    """
    # ----------------- 1. เตรียมตัวแปรวันที่และปฏิทิน -----------------
    target_dt = datetime(target_year, target_month, 1)
    prev_dt = target_dt - pd.DateOffset(months=1)
    
    num_days_curr = calendar.monthrange(target_year, target_month)[1]
    num_days_prev = calendar.monthrange(prev_dt.year, prev_dt.month)[1]
    
    days_cols = [str(d) for d in range(1, 32)]
    
    month_th = ["", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", 
                "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
    month_th_abbr = ["", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", 
                     "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
    
    title_text = f"ราคาน้ำมัน {month_th[target_month]} {target_year + 543}"
    col_avg_prev = f"avg {month_th_abbr[prev_dt.month]} {(prev_dt.year + 543) % 100}"
    col_avg_curr = f"avg {month_th_abbr[target_month]} {(target_year + 543) % 100}"
    
    df = raw_df.copy()
    df['Date'] = pd.to_datetime(df['Date'])

    # ----------------- 2. กำหนดโครงสร้าง 7 หมวดหลักและลิสต์สถานีบริการ -----------------
    priority_oil_configs = [
        {
            "priority": 1,
            "title": "1. เบนซิน ออกเทน 95 (ULG 95 RON)  (5220003)",
            "match": lambda name: ('เบนซิน' in name or 'ULG' in name.upper()) and ('95' in name) and ('แก๊ส' not in name and 'GAS' not in name.upper()),
            "brands": [
                'ปตท (PTT)',
                'เอสโซ่ (ESSO)',
                'เซฟรอน (CALTEX)',
                'ทีพีไอ (TPI)',
                'ภาคใต้เชื้อเพลิง (PT)',
                'สยามสหบริการ (SUSCO)',
                'คอสโม (COSMO)',
                'ปิโตนาส (PETRONAS)'
            ]
        },
        {
            "priority": 2,
            "title": "2. ดีเซลหมุนเร็ว (HSD, 0.005%S)  (5220004)",
            # ดึงจาก ดีเซลหมุนเร็ว HSD B7
            "match": lambda name: ('ดีเซล' in name or 'HSD' in name.upper()) and ('B7' in name.upper()),
            "brands": [
                'ปตท (PTT)',
                'เชลล์ (SHELL)',
                'เอสโซ่ (ESSO)',
                'เซฟรอน (CALTEX)',
                'TPI',
                'บางจาก (BCP)',
                'ภาคใต้เชื้อเพลิง (PT)',
                'สยามสหบริการ (SUSCO)',
                'ทั่วไป',
                'คอสโม (COSMO)',
                'เพียว (PURE)',
                'ปิโตนาส (PETRONAS)'
            ]
        },
        {
            "priority": 3,
            "title": "3. แก๊สโซฮอล ออกเทน 91 (Gasohol 91-E10)  (5220006)",
            "match": lambda name: ('91' in name) and ('แก๊ส' in name or 'GAS' in name.upper()),
            "brands": [
                'ปตท (PTT)',
                'บางจาก (BCP)',
                'เชลล์ (SHELL)',
                'เอสโซ่ (ESSO)',
                'คอสโม (COSMO)',
                'เซฟรอน (CALTEX)',
                'ทั่วไป',
                'ปิโตนาส (PETRONAS)',
                'ภาคใต้เชื้อเพลิง (PT)',
                'สยามสหบริการ (SUSCO)',
                'เพียว (PURE)',
                'TPI'
            ]
        },
        {
            "priority": 4,
            "title": "4. แก๊สโซฮอล ออกเทน 95 (Gasohol 95-E10)  (5220007)",
            "match": lambda name: ('95' in name) and ('แก๊ส' in name or 'GAS' in name.upper()) and ('E20' not in name.upper() and 'E85' not in name.upper()),
            "brands": [
                'ปตท (PTT)',
                'เชลล์ (SHELL)',
                'เอสโซ่ (ESSO)',
                'เซฟรอน (CALTEX)',
                'บางจาก (BCP)',
                'สยามสหบริการ (SUSCO)',
                'ภาคใต้เชื้อเพลิง (PT)',
                'เพียว (PURE)',
                'ปิโตนาส (PETRONAS)',
                'ทั่วไป',
                'คอสโม (COSMO)'
            ]
        },
        {
            "priority": 5,
            "title": "5. แก๊สโซฮอล E20 เบนซินไร้สารตะกั่ว 80% + แอลกอฮอล์ 20% (Gasohol 95-E20) (5220010)",
            "match": lambda name: 'E20' in name.upper(),
            "brands": [
                'ปตท (PTT)',
                'บางจาก (BCP)',
                'เซฟรอน (CALTEX)',
                'เชลล์ (SHELL)',
                'เอสโซ่ (ESSO)',
                'สยามสหบริการ (SUSCO)',
                'ภาคใต้เชื้อเพลิง (PT)'
            ]
        },
        {
            "priority": 6,
            "title": "6. แก๊สโซฮอล E85 ออกเทน 95 (Gasohol 95-E85) (5220011)",
            "match": lambda name: 'E85' in name.upper(),
            "brands": [
                'ปตท (PTT)',
                'บางจาก (BCP)'
            ]
        },
        {
            "priority": 7,
            "title": "7. ดีเซลหมุนเร็ว (B10) (HSD, 0.005%S)  (5220012)",
            "match": lambda name: 'B10' in name.upper(),
            "brands": [
                'ปตท (PTT)',
                'เชลล์ (SHELL)',
                'เอสโซ่ (ESSO)',
                'เซฟรอน (CALTEX)',
                'บางจาก (BCP)',
                'ภาคใต้เชื้อเพลิง (PT)',
                'สยามสหบริการ (SUSCO)'
            ]
        }
    ]

    # ฟังก์ชันค้นหาแบรนด์ในไฟล์ Raw ตามลำดับความสำคัญ (รวม Proxy/Fallback)
    def resolve_raw_brand(target_brand, available_raw_brands):
        t_clean = target_brand.upper().replace(' ', '')
        
        candidates = []
        if 'CALTEX' in t_clean or 'เซฟรอน' in t_clean or 'CHEVRON' in t_clean:
            candidates = ['CALTEX', 'CHEVRON', 'เซฟรอน', 'เชฟรอน']
        elif t_clean in ['TPI', 'ทีพีไอ', 'ทีพีไอ(TPI)', 'ทั่วไป']:
            candidates = [t_clean, 'PTT', 'ปตท']
        elif 'ภาคใต้' in t_clean or t_clean == 'PT' or 'คอสโม' in t_clean or 'COSMO' in t_clean:
            candidates = [t_clean, 'เอ็นเนอยี', 'ENERGY', 'PTG']
        elif 'ปิโตนาส' in t_clean or 'PETRONAS' in t_clean or 'สยามสห' in t_clean or 'SUSCO' in t_clean or 'SOSCO' in t_clean:
            candidates = [t_clean, 'SUSCO', 'สยามสหบริการ', 'สยามสห', 'SOSCO']
        elif 'ปตท' in t_clean or 'PTT' in t_clean:
            candidates = ['PTT', 'ปตท']
        elif 'เชลล์' in t_clean or 'SHELL' in t_clean:
            candidates = ['SHELL', 'เชลล์']
        elif 'เอสโซ่' in t_clean or 'ESSO' in t_clean:
            candidates = ['ESSO', 'เอสโซ่']
        elif 'บางจาก' in t_clean or 'BCP' in t_clean:
            candidates = ['BCP', 'บางจาก']
        elif 'เพียว' in t_clean or 'PURE' in t_clean:
            candidates = ['PURE', 'เพียว']
        else:
            candidates = [t_clean]

        for cand in candidates:
            cand_upper = cand.upper().replace(' ', '')
            for rb in available_raw_brands:
                rb_upper = str(rb).upper().replace(' ', '')
                if cand_upper in rb_upper or rb_upper in cand_upper:
                    return rb
        return None

    # จัดกลุ่มหมวดหมู่ผลิตภัณฑ์
    all_raw_products = [p for p in df['Product'].dropna().unique() if str(p).strip() != '']
    sorted_sections = []
    matched_raw_prods = set()

    for cfg in priority_oil_configs:
        matched_prod = None
        for p in all_raw_products:
            if p not in matched_raw_prods and cfg["match"](p):
                matched_prod = p
                matched_raw_prods.add(p)
                break
        sorted_sections.append({
            'raw_product': matched_prod,
            'title': cfg['title'],
            'fixed_brands': cfg['brands'],
            'is_priority': True
        })

    # หมวดหมู่น้ำมันอื่นๆ ที่นอกเหนือจาก 7 ข้อแรก (เรียงต่อท้าย และใส่ทุกสถานีบริการที่มีใน raw)
    other_idx = 8
    for p in all_raw_products:
        if p not in matched_raw_prods:
            sorted_sections.append({
                'raw_product': p,
                'title': f"{other_idx}. {p}",
                'fixed_brands': None,  # None = เอาทุกสถานีบริการที่มีใน raw
                'is_priority': False
            })
            other_idx += 1

    # ----------------- 3. ฟังก์ชันเตรียมข้อมูลตาราง -----------------
    def prepare_sheet_data(price_offset=0.0):
        rows = []
        
        for section in sorted_sections:
            raw_prod = section['raw_product']
            display_title = section['title']
            
            rows.append({
                'is_header': True,
                'title': display_title,
                'data': None
            })
            
            if raw_prod is not None:
                prod_df = df[df['Product'] == raw_prod].copy()
                if price_offset > 0:
                    prod_df['Price'] = prod_df['Price'] + price_offset

                curr_df = prod_df[(prod_df['Date'].dt.year == target_year) & (prod_df['Date'].dt.month == target_month)]
                prev_df = prod_df[(prod_df['Date'].dt.year == prev_dt.year) & (prod_df['Date'].dt.month == prev_dt.month)]
                
                prev_avg = prev_df[prev_df['Price'] > 0].groupby('Brand')['Price'].mean()
                prev_last_day_val = prev_df[prev_df['Date'].dt.day == num_days_prev].set_index('Brand')['Price']
                
                if not curr_df.empty:
                    pivot_curr = curr_df.pivot_table(index='Brand', columns=curr_df['Date'].dt.day, values='Price', aggfunc='last')
                else:
                    pivot_curr = pd.DataFrame()
                raw_brands_in_prod = list(prod_df['Brand'].dropna().unique())
            else:
                curr_df = pd.DataFrame()
                prev_avg = pd.Series(dtype=float)
                prev_last_day_val = pd.Series(dtype=float)
                pivot_curr = pd.DataFrame()
                raw_brands_in_prod = []

            # กำหนดรายชื่อสถานีบริการ
            if section['is_priority']:
                # 7 ข้อแรก: ล็อกตามลิสต์ที่กำหนด พร้อมทำ Fallback Lookup
                display_brands = [(b, resolve_raw_brand(b, raw_brands_in_prod)) for b in section['fixed_brands']]
            else:
                # ข้ออื่นๆ: แสดงครบทุกสถานีบริการที่มีในไฟล์ raw
                display_brands = [(b, b) for b in raw_brands_in_prod]

            for display_b_name, actual_raw_b in display_brands:
                r = {'Brand': display_b_name}
                
                if actual_raw_b is not None:
                    p_avg = prev_avg.get(actual_raw_b, np.nan)
                    r['avg_prev'] = p_avg if (pd.notna(p_avg) and p_avg > 0) else 0.00
                    
                    p_31 = prev_last_day_val.get(actual_raw_b, np.nan)
                    r['prev_31'] = p_31 if (pd.notna(p_31) and p_31 > 0) else 0.00
                    
                    valid_daily_prices = []
                    for d in range(1, 32):
                        if d <= num_days_curr and (actual_raw_b in pivot_curr.index) and (d in pivot_curr.columns):
                            val = pivot_curr.loc[actual_raw_b, d]
                            if pd.notna(val) and val > 0:
                                r[str(d)] = float(val)
                                valid_daily_prices.append(float(val))
                            else:
                                r[str(d)] = 0.00
                        else:
                            r[str(d)] = 0.00
                            
                    c_avg = np.mean(valid_daily_prices) if valid_daily_prices else 0.00
                    r['avg_curr'] = c_avg
                    
                    if r['avg_prev'] > 0 and c_avg > 0:
                        r['REL'] = (c_avg / r['avg_prev']) * 100.0
                    else:
                        r['REL'] = 0.00
                else:
                    # ถ้าหาไม่เจอให้ใส่ 0.00 ทั้งแถว
                    r['avg_prev'] = 0.00
                    r['prev_31'] = 0.00
                    for d in range(1, 32):
                        r[str(d)] = 0.00
                    r['avg_curr'] = 0.00
                    r['REL'] = 0.00
                
                rows.append({
                    'is_header': False,
                    'title': '',
                    'data': r
                })
                
        return rows

    # ----------------- 4. สร้างและจัดสไตล์ไฟล์ Excel -----------------
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    font_title = Font(name='Angsana New', size=16, bold=True)
    font_bold = Font(name='Angsana New', size=13, bold=True)
    font_regular = Font(name='Angsana New', size=12)
    
    fill_blue_header = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    fill_sub_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    fill_diff_up = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    font_diff_up = Font(name='Angsana New', size=12, color="C00000", bold=True)
    fill_diff_down = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    font_diff_down = Font(name='Angsana New', size=12, color="375623", bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='A6A6A6'),
        right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'),
        bottom=Side(style='thin', color='A6A6A6')
    )
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    # 4.1 Sheet 1 และ Sheet 2
    sheets_config = [
        ('ราคาน้ำมันขายปลีก', 0.0),
        ('ราคารวมหัวจ่าย (+0.05)', 0.05)
    ]
    
    headers_main = ['ปั๊มน้ำมัน/วันที่', col_avg_prev, '31'] + days_cols + [col_avg_curr, 'REL']

    for sheet_name, offset in sheets_config:
        ws = wb.create_sheet(title=sheet_name)
        matrix = prepare_sheet_data(price_offset=offset)
        
        ws.cell(row=1, column=1, value=title_text).font = font_title
        
        for col_idx, h_text in enumerate(headers_main, 1):
            cell = ws.cell(row=3, column=col_idx, value=h_text)
            cell.font = font_bold
            cell.alignment = align_center
            cell.border = thin_border
            if col_idx in [2, len(headers_main) - 1]:
                cell.fill = fill_yellow
            elif col_idx == 3:
                cell.fill = fill_sub_green

        current_row = 4
        for item in matrix:
            if item['is_header']:
                c = ws.cell(row=current_row, column=1, value=item['title'])
                c.font = font_bold
                c.alignment = align_left
                for col_idx in range(1, len(headers_main) + 1):
                    ws.cell(row=current_row, column=col_idx).fill = fill_blue_header
                    ws.cell(row=current_row, column=col_idx).border = thin_border
                current_row += 1
            else:
                d = item['data']
                b_cell = ws.cell(row=current_row, column=1, value=d['Brand'])
                b_cell.font = font_regular
                b_cell.alignment = align_left
                b_cell.border = thin_border
                
                p_cell = ws.cell(row=current_row, column=2, value=d['avg_prev'])
                p_cell.number_format = '0.00'
                p_cell.font = font_regular
                p_cell.fill = fill_yellow
                p_cell.alignment = align_right
                p_cell.border = thin_border
                
                prev31_cell = ws.cell(row=current_row, column=3, value=d['prev_31'])
                prev31_cell.number_format = '0.00'
                prev31_cell.font = font_regular
                prev31_cell.fill = fill_sub_green
                prev31_cell.alignment = align_right
                prev31_cell.border = thin_border
                
                for idx, day_str in enumerate(days_cols, 4):
                    day_cell = ws.cell(row=current_row, column=idx, value=d[day_str])
                    day_cell.number_format = '0.00'
                    day_cell.font = font_regular
                    day_cell.alignment = align_right
                    day_cell.border = thin_border
                
                c_avg_cell = ws.cell(row=current_row, column=len(headers_main) - 1, value=d['avg_curr'])
                c_avg_cell.number_format = '0.00'
                c_avg_cell.font = font_regular
                c_avg_cell.fill = fill_yellow
                c_avg_cell.alignment = align_right
                c_avg_cell.border = thin_border
                
                rel_cell = ws.cell(row=current_row, column=len(headers_main), value=d['REL'])
                rel_cell.number_format = '0.00'
                rel_cell.font = font_regular
                rel_cell.alignment = align_right
                rel_cell.border = thin_border
                
                current_row += 1

    # 4.2 Sheet 3 (ผลต่างราคารายวันเทียบวันก่อนหน้า)
    ws3 = wb.create_sheet(title='ผลต่างราคารายวัน')
    ws3.cell(row=1, column=1, value=f"{title_text} (ส่วนต่างราคาเทียบวันก่อนหน้า)").font = font_title
    
    headers_diff = ['ปั๊มน้ำมัน/วันที่', '31'] + days_cols + ['AI (Avg Diff)']
    
    for col_idx, h_text in enumerate(headers_diff, 1):
        cell = ws3.cell(row=3, column=col_idx, value=h_text)
        cell.font = font_bold
        cell.alignment = align_center
        cell.border = thin_border
        if col_idx == 2:
            cell.fill = fill_sub_green
        elif col_idx == len(headers_diff):
            cell.fill = fill_yellow

    base_data = prepare_sheet_data(price_offset=0.0)
    current_row = 4
    
    for item in base_data:
        if item['is_header']:
            c = ws3.cell(row=current_row, column=1, value=item['title'])
            c.font = font_bold
            c.alignment = align_left
            for col_idx in range(1, len(headers_diff) + 1):
                ws3.cell(row=current_row, column=col_idx).fill = fill_blue_header
                ws3.cell(row=current_row, column=col_idx).border = thin_border
            current_row += 1
        else:
            d = item['data']
            b_cell = ws3.cell(row=current_row, column=1, value=d['Brand'])
            b_cell.font = font_regular
            b_cell.alignment = align_left
            b_cell.border = thin_border
            
            cell_prev31 = ws3.cell(row=current_row, column=2, value=0.00)
            cell_prev31.number_format = '0.00'
            cell_prev31.fill = fill_sub_green
            cell_prev31.border = thin_border
            cell_prev31.alignment = align_right
            
            diff_values = []
            prev_price = d['prev_31'] if (pd.notna(d['prev_31']) and d['prev_31'] > 0) else 0.00
            
            for idx, day_str in enumerate(days_cols, 3):
                curr_price = d[day_str]
                
                if curr_price > 0 and prev_price > 0:
                    diff = round(curr_price - prev_price, 2)
                    diff_values.append(diff)
                else:
                    diff = 0.00
                
                cell_diff = ws3.cell(row=current_row, column=idx, value=diff)
                cell_diff.number_format = '0.00'
                cell_diff.alignment = align_right
                cell_diff.border = thin_border
                
                if diff > 0:
                    cell_diff.fill = fill_diff_up
                    cell_diff.font = font_diff_up
                elif diff < 0:
                    cell_diff.fill = fill_diff_down
                    cell_diff.font = font_diff_down
                else:
                    cell_diff.font = font_regular
                
                if curr_price > 0:
                    prev_price = curr_price
            
            avg_diff = np.mean(diff_values) if diff_values else 0.00
            ai_cell = ws3.cell(row=current_row, column=len(headers_diff), value=avg_diff)
            ai_cell.number_format = '0.00'
            ai_cell.font = font_bold
            ai_cell.fill = fill_yellow
            ai_cell.alignment = align_right
            ai_cell.border = thin_border
            
            current_row += 1

    # ----------------- 5. ตั้งค่า Freeze Panes และความกว้างคอลัมน์ -----------------
    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.column_dimensions['A'].width = 24
        for col_letter in [get_column_letter(c) for c in range(2, sheet.max_column + 1)]:
            sheet.column_dimensions[col_letter].width = 7.5

    wb.save(output_path)
    print(f"สร้างไฟล์รายงานสำเร็จเรียบร้อยที่: {output_path}")
